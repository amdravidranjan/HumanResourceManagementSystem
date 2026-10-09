"""Builds docs/slides/HRMS_Demo.pptx (16:9, 12 slides) from evaluation/results and evaluation/graphs.

Usage: python docs/slides/build_slides.py
"""
import json
import pathlib

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = pathlib.Path(__file__).resolve().parents[2]
RES = ROOT / "evaluation" / "results"
GR = ROOT / "evaluation" / "graphs"
OUT = pathlib.Path(__file__).resolve().parent / "HRMS_Demo.pptx"

NAVY = RGBColor(0x1F, 0x3A, 0x8A)
BLUE = RGBColor(0x2F, 0x5B, 0xEA)
INK = RGBColor(0x1D, 0x24, 0x33)
MUTED = RGBColor(0x5D, 0x6B, 0x82)
LIGHT = RGBColor(0xEE, 0xF2, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEAM = [("Dravid Ranjan A.M", "3122245001046"), ("Guhanesh M.L", "3122245001049"),
        ("Hansika N.M", "3122245001052"), ("Saisaravanan S", "3122245001310")]


def j(name):
    return json.loads((RES / f"{name}.json").read_text(encoding="utf-8"))


P, B, RUN1 = j("components_proposed"), j("components_baseline"), j("components_proposed_run1")
L = {k: j(f"load_{k}") for k in ["low", "medium", "high", "baseline_low", "baseline_medium", "baseline_high",
                                 "node_failure", "shard_failure"]}
SHARD = j("scenario_add_shard")["result"]
CRAWL = RUN1["stats"]["crawler"]
mongo_p = sum(L[k]["resources"]["mongo"]["avgCpuPercentTotal"] for k in ["low", "medium", "high"]) / 3
mongo_b = sum(L[k]["resources"]["mongo"]["avgCpuPercentTotal"] for k in ["baseline_low", "baseline_medium",
                                                                          "baseline_high"]) / 3
trie_x = P["autocomplete"]["regexScatter"]["avgMs"] / P["autocomplete"]["trie"]["avgMs"]

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = 13.333, 7.5
N = [0]


def box(slide, x, y, w, h, fill=None, line=None):
    s = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, lines, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, bullet=False, gap=6):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, ln in enumerate([lines] if isinstance(lines, str) else lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(gap)
        parts = ln if isinstance(ln, tuple) else (ln,)
        if bullet:
            r = p.add_run()
            r.text = "•  "
            r.font.size, r.font.color.rgb = Pt(size), BLUE
        for k, part in enumerate(parts):
            r = p.add_run()
            r.text = part
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.bold = bold or (len(parts) > 1 and k == 0)
    return tb


def slide(title, kicker=None):
    s = prs.slides.add_slide(BLANK)
    N[0] += 1
    box(s, 0, 0, W, 0.12, fill=NAVY)
    if kicker:
        text(s, 0.6, 0.35, 12, 0.4, kicker.upper(), size=12, color=BLUE, bold=True)
    text(s, 0.6, 0.62, 12.2, 0.8, title, size=30, color=NAVY, bold=True)
    text(s, 11.9, 7.05, 1.1, 0.35, str(N[0]), size=11, color=MUTED, align=PP_ALIGN.RIGHT)
    text(s, 0.6, 7.05, 8, 0.35, "Scalable HRM System  |  UCS3513 System Design Lab", size=11, color=MUTED)
    return s


def image(s, path, x, y, w, h):
    """Place an image inside the box (x, y, w, h), keeping its aspect ratio and centring it."""
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    dw, dh = iw * scale, ih * scale
    s.shapes.add_picture(str(path), Inches(x + (w - dw) / 2), Inches(y + (h - dh) / 2), Inches(dw), Inches(dh))


def stat(s, x, y, w, big, small, color=BLUE):
    box(s, x, y, w, 1.35, fill=LIGHT)
    text(s, x + 0.15, y + 0.1, w - 0.3, 0.7, big, size=28, color=color, bold=True)
    text(s, x + 0.15, y + 0.75, w - 0.3, 0.6, small, size=12, color=MUTED)


def table(s, x, y, w, header, rows, col_w, size=12, row_h=0.42):
    shp = s.shapes.add_table(len(rows) + 1, len(header), Inches(x), Inches(y), Inches(w), Inches(row_h * (len(rows) + 1)))
    t = shp.table
    for i, cw in enumerate(col_w):
        t.columns[i].width = Inches(cw)
    for r, row in enumerate([header] + rows):
        for c, v in enumerate(row):
            cell = t.cell(r, c)
            cell.text = str(v)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            p = cell.text_frame.paragraphs[0]
            p.runs[0].font.size = Pt(size)
            p.runs[0].font.bold = r == 0
            p.runs[0].font.color.rgb = WHITE if r == 0 else INK
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (LIGHT if r % 2 == 0 else WHITE)
    return t


# 1 ---------------------------------------------------------------- title
s = prs.slides.add_slide(BLANK)
N[0] += 1
box(s, 0, 0, W, H, fill=NAVY)
box(s, 0.6, 1.2, 0.12, 2.2, fill=BLUE)
text(s, 0.95, 1.1, 11.5, 0.5, "UCS3513 SYSTEM DESIGN LABORATORY  |  MINI PROJECT  |  PROBLEM STATEMENT 6", size=14,
     color=RGBColor(0xB8, 0xC7, 0xF5), bold=True)
text(s, 0.95, 1.6, 12.0, 1.0, "Scalable Human Resource Management System", size=38, color=WHITE, bold=True)
text(s, 0.95, 2.8, 11.5, 0.6, "Gateway + 3 HR nodes + 4 MongoDB shards + Redis, with 8 system-design components",
     size=18, color=RGBColor(0xDC, 0xE4, 0xF7))
for i, (n, r) in enumerate(TEAM):
    x = 0.95 + i * 3.0
    box(s, x, 4.5, 2.8, 1.1, fill=RGBColor(0x2A, 0x4A, 0xA5))
    text(s, x + 0.15, 4.6, 2.5, 0.45, n, size=16, color=WHITE, bold=True)
    text(s, x + 0.15, 5.05, 2.5, 0.4, r, size=14, color=RGBColor(0xB8, 0xC7, 0xF5))
text(s, 0.95, 6.4, 11.5, 0.5, "SSN College of Engineering  |  B.E. CSE, V Semester  |  October 2026", size=14,
     color=RGBColor(0xB8, 0xC7, 0xF5))

# 2 ---------------------------------------------------------------- problem
s = slide("An HRMS that keeps working as the organisation grows", "Problem and goals")
text(s, 0.6, 1.7, 6.0, 4.8, [
    ("Workload  ", "read-heavy lookups, daily check-in bursts, month-end payroll"),
    ("Growth  ", "attendance grows as employees x days"),
    ("Concurrency  ", "several servers create records at the same time"),
    ("Outside data  ", "job listings from external portals"),
    ("Availability  ", "must survive a server or a database failing"),
], size=17, bullet=True, gap=14)
box(s, 7.0, 1.7, 5.7, 4.8, fill=LIGHT)
text(s, 7.3, 1.9, 5.2, 0.5, "What we built", size=18, color=NAVY, bold=True)
text(s, 7.3, 2.5, 5.2, 4.0, [
    "Employees, attendance, leave, payroll, recruitment, performance, notifications",
    "3 roles: HR admin, manager, employee",
    "10,000 synthetic employees across 4 shards",
    "Evaluated at 10 / 50 / 200 users against a single-server baseline",
    "Node and shard failure, node and shard scaling",
], size=15, bullet=True, gap=10)

# 3 ---------------------------------------------------------------- architecture
s = slide("Architecture", "Deliverable 2")
image(s, ROOT / "docs" / "architecture.png", 0.6, 1.5, 12.1, 5.45)

# 4 ---------------------------------------------------------------- component mapping
s = slide("Eight components, each with a real HR job", "Deliverable 1")
table(s, 0.6, 1.55, 12.1, ["Component", "Used for", "Design"], [
    ("Unique ID generator", "IDs for every record, from 3 nodes at once", "Snowflake: 41-bit time, 10-bit node, 12-bit seq"),
    ("Sharding", "All employee-owned data co-located by employeeId", "App-level hash sharding over 4 MongoDB"),
    ("Consistent hashing", "Gateway -> node, employeeId -> shard", "MurmurHash3 ring, 150 virtual nodes"),
    ("Key-value store", "Sessions, cache, limiter, short links, shard map", "Redis cache-aside + TTL, pub/sub"),
    ("Rate limiter", "Login brute force, payroll runs, all APIs", "Token bucket, atomic Redis Lua"),
    ("Autocomplete", "Employee and job search as you type", "Trie with cached top-10 per prefix"),
    ("URL shortener", "Public job links, expiring payslip links", "Base62(Snowflake), Redis TTL, 302"),
    ("Web crawler", "Import jobs from 3 portals", "BFS, robots.txt, 200 ms politeness, dedup"),
], [2.6, 4.9, 4.6], size=13, row_h=0.6)

# 5 ---------------------------------------------------------------- request flow & sharding
s = slide("One request, one node, one shard", "Request flow and sharding")
text(s, 0.6, 1.6, 6.2, 5.2, [
    "Browser -> gateway :8080 with session token",
    "Rate limiter: token bucket in Redis (429 if empty)",
    "Node ring picks the HR node for this employee; retry on next node if it fails",
    "Node checks session and role (Redis)",
    "Cache hit -> answer from Redis; miss -> shard ring -> MongoDB -> cache",
    "All of an employee's data is on one shard, so most requests touch one shard",
], size=16, bullet=True, gap=12)
image(s, GR / "shard_distribution.png", 7.0, 1.6, 5.8, 4.2)
text(s, 7.0, 5.9, 5.8, 0.9, f"Adding shard-3 moved {SHARD['employeesMovedPercent']:.1f}% of employees "
                            f"({SHARD['documentsMoved']:,} documents) in {SHARD['durationMs'] / 1000:.1f} s",
     size=14, color=MUTED, align=PP_ALIGN.CENTER)

# 6 ---------------------------------------------------------------- snowflake + shortener
s = slide("Unique IDs and short links", "Snowflake + URL shortener")
sn = P["snowflake"]
stat(s, 0.6, 1.7, 3.8, f"{sn['singleGeneratorIdsPerSecond'] / 1e6:.1f} M/s", "IDs per second, one generator")
stat(s, 4.75, 1.7, 3.8, f"{sn['distinctNodeIds']['collisions']}", "collisions in 300k IDs from 3 nodes")
stat(s, 8.9, 1.7, 3.8, f"{sn['sameNodeIdMisconfigured']['collisions']:,}", "collisions if all nodes share an ID",
     color=RGBColor(0xC8, 0x2E, 0x3A))
text(s, 0.6, 3.4, 12, 0.5, "id = (ms since 2026-01-01) << 22  |  nodeId << 12  |  sequence", size=18, color=NAVY,
     bold=True)
rd = P["redirect"]
text(s, 0.6, 4.2, 12, 2.6, [
    "Short code = Base62 of a Snowflake ID, served at /s/{code}",
    "Redis holds code -> URL with a TTL (payslip 24 h, offer 7 d); MongoDB keeps a durable copy",
    f"Redirect: {rd['avgMs']:.1f} ms average, p95 {rd['p95Ms']:.1f} ms, {rd['redirects302']}/{rd['requests']} correct 302s",
], size=17, bullet=True, gap=12)

# 7 ---------------------------------------------------------------- hashing
s = slide("Consistent hashing: even spread, minimal movement", "Consistent hashing + rebalancing")
image(s, GR / "hashing.png", 0.6, 1.5, 8.2, 5.3)
h = P["hashing"]
stat(s, 9.1, 1.7, 3.6, f"{h['addFourthShard']['consistentMovedPercent']:.1f}%", "keys moved 3 -> 4 shards")
stat(s, 9.1, 3.25, 3.6, f"{h['addFourthShard']['moduloMovedPercent']:.1f}%", "with hash mod N",
     color=RGBColor(0xC8, 0x2E, 0x3A))
stat(s, 9.1, 4.8, 3.6, f"{h['distributionByVirtualNodes'][3]['stdDevPercent']:.1f}%",
     "std-dev of load, 150 virtual nodes")

# 8 ---------------------------------------------------------------- kv + autocomplete
s = slide("Redis and the trie take the reads off MongoDB", "Key-value store + autocomplete")
image(s, GR / "kv_latency.png", 0.6, 1.5, 6.0, 3.9)
image(s, GR / "autocomplete_latency.png", 6.8, 1.5, 6.0, 3.9)
text(s, 0.6, 5.55, 12.1, 1.4, [
    f"Cache hit ratio {RUN1['kv']['nodeCacheHitRatio'] * 100:.1f}% under live traffic; MongoDB CPU {mongo_b:.0f}% "
    f"(baseline) -> {mongo_p:.0f}% (proposed) at the same request rate",
    f"Trie {P['autocomplete']['trie']['avgMs']:.3f} ms vs regex across shards "
    f"{P['autocomplete']['regexScatter']['avgMs']:.1f} ms (over {int(trie_x // 1000) * 1000:,}x), precision@10 = "
    f"{P['autocomplete']['precisionAt10']:.2f}",
], size=15, bullet=True, gap=8)

# 9 ---------------------------------------------------------------- limiter + crawler
s = slide("Rate limiter and web crawler", "Rate limiter + web crawler")
image(s, GR / "rate_limiter.png", 0.6, 1.5, 7.6, 3.9)
rl, bf = P["rateLimiter"], P["loginBruteForce"]
text(s, 0.6, 5.55, 7.6, 1.4, [
    f"Burst of {rl['sent']}: {rl['allowed']} allowed (theory ~{rl['expectedAllowedApprox']}), "
    f"{rl['rejectionRate'] * 100:.1f}% rejected with 429",
    f"Login brute force: {bf['blocked429']} of {bf['attempts']} attempts blocked after 5 failures",
], size=14, bullet=True, gap=8)
box(s, 8.6, 1.5, 4.1, 5.3, fill=LIGHT)
text(s, 8.85, 1.65, 3.7, 0.5, "Crawler (3 portals)", size=18, color=NAVY, bold=True)
text(s, 8.85, 2.3, 3.7, 4.4, [
    ("Pages  ", f"{CRAWL['pagesFetched']} at {CRAWL['pagesPerSecond']}/s"),
    ("Jobs  ", f"{CRAWL['jobsFound']} unique"),
    ("Duplicates  ", f"{CRAWL['duplicatesDropped']} dropped"),
    ("robots.txt  ", f"{CRAWL['blockedByRobots']} pages skipped"),
    ("Failures  ", f"{CRAWL['pagesFailed']}"),
    "BFS, depth 3, 200 ms politeness per host",
], size=15, bullet=True, gap=12)

# 10 ---------------------------------------------------------------- load
s = slide("Load test: proposed vs baseline", "Deliverable 4")
image(s, GR / "latency_by_load.png", 0.6, 1.45, 7.4, 3.0)
rows = []
for lvl, a, b in [("10 users", "low", "baseline_low"), ("50 users", "medium", "baseline_medium"),
                  ("200 users", "high", "baseline_high")]:
    rows.append((lvl, f"{L[a]['throughputRps']:.0f} / {L[b]['throughputRps']:.0f}",
                 f"{L[a]['p95Ms']:.0f} / {L[b]['p95Ms']:.0f}", f"{L[a]['successRate'] * 100:.0f}%",
                 f"{L[a]['resources']['mongo']['avgCpuPercentTotal']:.0f} / "
                 f"{L[b]['resources']['mongo']['avgCpuPercentTotal']:.0f}"))
table(s, 8.2, 1.55, 4.5, ["Load", "Req/s", "p95 ms", "OK", "Mongo CPU %"], rows, [1.0, 0.95, 0.95, 0.6, 1.0],
      size=11, row_h=0.45)
text(s, 8.2, 3.5, 4.5, 0.5, "proposed / baseline", size=11, color=MUTED)
text(s, 0.6, 4.7, 12.1, 2.2, [
    "100% success at every level, for both systems",
    "Throughput is flat (~360-395 req/s): all 10 containers and the load generator share one laptop CPU; "
    "latency grows with users as Little's law predicts",
    f"The real gain is database headroom: MongoDB CPU about {mongo_b / mongo_p:.0f}x lower thanks to the cache "
    "and the trie",
], size=15, bullet=True, gap=8)

# 11 ---------------------------------------------------------------- failures
s = slide("Failure scenarios", "Deliverable 5")
image(s, GR / "node_failure_timeline.png", 0.6, 1.45, 6.0, 2.9)
image(s, GR / "shard_failure_timeline.png", 6.75, 1.45, 6.0, 2.9)
nf, sf = L["node_failure"], L["shard_failure"]
text(s, 0.6, 4.5, 6.0, 2.4, [
    ("HR node killed  ", f"{nf['successRate'] * 100:.0f}% success, p95 {nf['p95Ms']:.0f} ms"),
    "Gateway retries on the next node; health check removes it in 2 s; sessions in Redis",
], size=14, bullet=True, gap=8)
text(s, 6.75, 4.5, 6.0, 2.4, [
    ("Shard killed  ", f"{sf['successRate'] * 100:.2f}% success"),
    "503 only for users whose data is on that shard (check-in, payslips, notifications)",
    "Cached reads and autocomplete never failed; recovered on restart",
], size=14, bullet=True, gap=8)

# 12 ---------------------------------------------------------------- trade-offs + demo
s = slide("Trade-offs, limitations and live demo", "Wrap-up")
text(s, 0.6, 1.6, 6.0, 5.2, [
    ("App-level sharding  ", "full control, but the app routes every query"),
    ("Cache-aside  ", "less DB load, stale for up to the TTL"),
    ("Redis  ", "shared state makes nodes stateless, but is one point of failure"),
    ("Gateway retry  ", "hides node failures, but needs idempotent requests"),
    ("Limitations  ", "one machine; no replica sets; synthetic data only"),
], size=15, bullet=True, gap=12)
box(s, 7.0, 1.6, 5.7, 5.2, fill=LIGHT)
text(s, 7.3, 1.75, 5.2, 0.5, "Live demo", size=18, color=NAVY, bold=True)
text(s, 7.3, 2.35, 5.2, 4.4, [
    "Admin dashboard, shards and serving node",
    "Employee search with autocomplete",
    "Leave apply -> manager approve -> balance",
    "Payroll run, payslip short link",
    "Crawl job portals, public apply",
    "docker stop hrms-node-2: app keeps working",
    "Add node 4; add shard-3 and rebalance",
], size=15, bullet=True, gap=8)

prs.save(OUT)
print("wrote", OUT, "-", N[0], "slides")
